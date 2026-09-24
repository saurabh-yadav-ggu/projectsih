from pathlib import Path
from typing import Tuple, List, Dict, Any

from app.verification.artifact import Artifact


def verify_docx(file_path: Path) -> Tuple[bool, List[str], List[str]]:
    checks = ["File exists and is non-empty"]
    errors = []
    try:
        from docx import Document
        doc = Document(str(file_path))
        p_count = len(doc.paragraphs)
        t_count = len(doc.tables)
        checks.append(f"Parsed {p_count} paragraphs and {t_count} tables")
        if p_count == 0 and t_count == 0:
            errors.append("DOCX contains neither paragraphs nor tables.")
    except Exception as e:
        errors.append(f"Failed to open/parse DOCX: {str(e)}")
    return (len(errors) == 0), checks, errors


def verify_xlsx(file_path: Path) -> Tuple[bool, List[str], List[str]]:
    checks = ["File exists and is non-empty"]
    errors = []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(file_path), data_only=False)
        sheet_names = wb.sheetnames
        checks.append(f"Parsed sheets: {', '.join(sheet_names)}")
        if not sheet_names:
            errors.append("XLSX has no sheets.")
        else:
            ws = wb.active
            if ws.max_row is None or ws.max_row < 1:
                errors.append("Active sheet has 0 rows.")
            else:
                checks.append(f"Active sheet '{ws.title}' has {ws.max_row} rows and {ws.max_column} columns")
    except Exception as e:
        errors.append(f"Failed to open/parse XLSX: {str(e)}")
    return (len(errors) == 0), checks, errors


def verify_pptx(file_path: Path) -> Tuple[bool, List[str], List[str]]:
    checks = ["File exists and is non-empty"]
    errors = []
    try:
        from pptx import Presentation
        prs = Presentation(str(file_path))
        slide_count = len(prs.slides)
        checks.append(f"Parsed presentation with {slide_count} slides")
        if slide_count == 0:
            errors.append("Presentation has 0 slides.")
    except Exception as e:
        errors.append(f"Failed to open/parse PPTX: {str(e)}")
    return (len(errors) == 0), checks, errors


def verify_pdf(file_path: Path) -> Tuple[bool, List[str], List[str]]:
    checks = ["File exists and is non-empty"]
    errors = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        page_count = len(reader.pages)
        checks.append(f"Parsed PDF with {page_count} pages")
        if page_count == 0:
            errors.append("PDF has 0 pages.")
    except Exception as e:
        errors.append(f"Failed to open/parse PDF: {str(e)}")
    return (len(errors) == 0), checks, errors


def verify_html(file_path: Path) -> Tuple[bool, List[str], List[str]]:
    checks = ["File exists and is non-empty"]
    errors = []
    try:
        content = file_path.read_text(encoding="utf-8", errors="replace")
        checks.append(f"Read {len(content)} characters of HTML content")
        lower = content.lower()
        if "<html" not in lower or "<body" not in lower:
            errors.append("HTML document missing basic <html> or <body> tags.")
    except Exception as e:
        errors.append(f"Failed to read HTML file: {str(e)}")
    return (len(errors) == 0), checks, errors


def verify_artifact(artifact: Artifact) -> Artifact:
    """
    Programmatically verifies an artifact based on its file extension.
    Updates verification_status, checks, and errors.
    """
    path = Path(artifact.path)
    if not path.exists():
        artifact.verification_status = "failed"
        artifact.verification_errors = [f"Artifact file not found at {artifact.path}"]
        return artifact

    if path.stat().st_size == 0:
        artifact.verification_status = "failed"
        artifact.verification_errors = ["Artifact file is 0 bytes (empty)."]
        return artifact

    artifact.size = path.stat().st_size
    ext = artifact.file_type.lower()

    if ext == "docx":
        passed, checks, errors = verify_docx(path)
    elif ext == "xlsx":
        passed, checks, errors = verify_xlsx(path)
    elif ext == "pptx":
        passed, checks, errors = verify_pptx(path)
    elif ext == "pdf":
        passed, checks, errors = verify_pdf(path)
    elif ext in ["html", "htm"]:
        passed, checks, errors = verify_html(path)
    elif ext in ["csv", "txt", "md", "json"]:
        passed, checks, errors = True, ["File exists and contains text data"], []
    else:
        passed, checks, errors = True, ["Generic file exists and is non-empty"], []

    artifact.verification_status = "passed" if passed else "failed"
    artifact.verification_checks = checks
    artifact.verification_errors = errors
    return artifact
