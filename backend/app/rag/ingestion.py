from pathlib import Path
from typing import List, Optional
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader

from app.rag.embeddings import text_splitter, vector_store


def load_pdf(file_path: str) -> List[Document]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF not found: {file_path}")
    if path.suffix.lower() != ".pdf":
        raise ValueError("File must be a PDF.")

    loader = PyPDFLoader(str(path))
    return loader.load()


def load_text_file(file_path: str) -> List[Document]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    text = path.read_text(encoding="utf-8", errors="replace")
    return [
        Document(
            page_content=text,
            metadata={"source": str(path)}
        )
    ]


def add_documents(
    documents: List[Document],
    thread_id: Optional[str] = None,
    user_id: Optional[int] = None
) -> int:
    chunks = text_splitter.split_documents(documents)
    if not chunks:
        return 0

    for chunk in chunks:
        if thread_id:
            chunk.metadata["thread_id"] = str(thread_id)
        if user_id:
            chunk.metadata["user_id"] = user_id

    vector_store.add_documents(documents=chunks)
    return len(chunks)


def add_pdf(file_path: str, thread_id: Optional[str] = None, user_id: Optional[int] = None) -> int:
    documents = load_pdf(file_path)
    filename = Path(file_path).name
    for doc in documents:
        doc.metadata["file_type"] = "pdf"
        doc.metadata["filename"] = filename

    return add_documents(documents, thread_id=thread_id, user_id=user_id)


def add_text_file(file_path: str, thread_id: Optional[str] = None, user_id: Optional[int] = None) -> int:
    documents = load_text_file(file_path)
    filename = Path(file_path).name
    for doc in documents:
        doc.metadata["file_type"] = "text"
        doc.metadata["filename"] = filename

    return add_documents(documents, thread_id=thread_id, user_id=user_id)


def load_docx(file_path: str) -> List[Document]:
    import docx
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")
    doc = docx.Document(str(path))
    text_parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            text_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    full_text = "\n".join(text_parts)
    return [Document(page_content=full_text, metadata={"source": str(path), "filename": path.name, "file_type": "docx"})]


def load_csv(file_path: str) -> List[Document]:
    import csv
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"CSV not found: {file_path}")
    lines = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            if row:
                lines.append(", ".join(row))
    return [Document(page_content="\n".join(lines), metadata={"source": str(path), "filename": path.name, "file_type": "csv"})]


def add_docx(file_path: str, thread_id: Optional[str] = None, user_id: Optional[int] = None) -> int:
    documents = load_docx(file_path)
    return add_documents(documents, thread_id=thread_id, user_id=user_id)


def add_csv(file_path: str, thread_id: Optional[str] = None, user_id: Optional[int] = None) -> int:
    documents = load_csv(file_path)
    return add_documents(documents, thread_id=thread_id, user_id=user_id)


def add_file(file_path: str, thread_id: Optional[str] = None, user_id: Optional[int] = None) -> int:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return add_pdf(file_path, thread_id=thread_id, user_id=user_id)
    elif ext in [".txt", ".md"]:
        return add_text_file(file_path, thread_id=thread_id, user_id=user_id)
    elif ext == ".docx":
        return add_docx(file_path, thread_id=thread_id, user_id=user_id)
    elif ext == ".csv":
        return add_csv(file_path, thread_id=thread_id, user_id=user_id)
    return 0
