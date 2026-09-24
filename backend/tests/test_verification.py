import pytest
from pathlib import Path
from docx import Document
from openpyxl import Workbook

from app.verification.artifact import Artifact
from app.verification.document import verify_artifact, verify_docx, verify_xlsx, verify_html


def test_verify_docx_valid(tmp_path: Path):
    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("This is a verified paragraph.")
    doc.save(doc_path)

    art = Artifact.from_path(str(doc_path))
    verified = verify_artifact(art)
    assert verified.verification_status == "passed"
    assert len(verified.verification_errors) == 0


def test_verify_xlsx_valid(tmp_path: Path):
    xlsx_path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Category", "Value"])
    ws.append(["Revenue", 50000])
    wb.save(xlsx_path)

    art = Artifact.from_path(str(xlsx_path))
    verified = verify_artifact(art)
    assert verified.verification_status == "passed"


def test_verify_html_valid(tmp_path: Path):
    html_path = tmp_path / "report.html"
    html_path.write_text("<!DOCTYPE html><html><head><title>Report</title></head><body><h1>Report</h1></body></html>", encoding="utf-8")

    art = Artifact.from_path(str(html_path))
    verified = verify_artifact(art)
    assert verified.verification_status == "passed"


def test_verify_nonexistent_fails():
    art = Artifact.from_path("/non/existent/path/report.docx")
    verified = verify_artifact(art)
    assert verified.verification_status == "failed"
