import pytest
import os
from pathlib import Path
from docx import Document
import openpyxl
from pptx import Presentation

from app.agent.content_planner import (
    plan_and_generate_content,
    _validate_and_clean_content,
    _build_deterministic_grounded_content,
    _sanitize_string,
)
from app.agent.subagents.document_agent import (
    _build_fallback_document,
    _suggest_filename,
    run_document_agent,
)


def test_sanitize_string():
    raw = "Here is a section [Insert text here] with lorem ipsum dolor sit amet. Done."
    cleaned = _sanitize_string(raw)
    assert "[Insert text here]" not in cleaned
    assert "lorem ipsum" not in cleaned.lower()
    assert "Here is a section" in cleaned


def test_validate_and_clean_content_generic_title_override():
    raw_content = {
        "title": "Business Report",
        "topic": "Amazon KDP Opportunities in India",
        "sections": [
            {
                "heading": "1. Market Overview",
                "paragraphs": ["eBook publishing in India is growing rapidly at 25% CAGR."],
                "bullet_points": ["Vernacular languages show highest demand."],
            }
        ]
    }
    cleaned = _validate_and_clean_content(raw_content, query="Amazon KDP business opportunities in India", doc_format="docx")
    assert cleaned["title"] != "Business Report"
    assert "Amazon Kdp" in cleaned["title"]
    assert len(cleaned["sections"]) == 1
    assert "25% CAGR" in cleaned["sections"][0]["paragraphs"][0]


def test_build_deterministic_grounded_content():
    query = "Create a report on Amazon KDP business opportunities in India"
    content = _build_deterministic_grounded_content(query, doc_format="docx")
    assert "Amazon Kdp" in content["title"]
    assert len(content["sections"]) >= 3
    # Verify no employee review keywords leaked
    assert "employee" not in content["title"].lower()
    assert "leave" not in content["title"].lower()


def test_build_fallback_document_docx(tmp_path: Path):
    target_docx = tmp_path / "kdp_report.docx"
    content_plan = {
        "title": "Amazon KDP Business Opportunities in India",
        "subtitle": "Market Analysis and Financial Forecasts",
        "sections": [
            {
                "heading": "1. Market Sizing and Vernacular Demand",
                "paragraphs": ["India's digital reading market is expanding across Hindi and regional languages."],
                "bullet_points": ["70% royalty tier available on select pricing thresholds."],
                "table": {
                    "headers": ["Metric", "Value", "Strategic Implication"],
                    "rows": [
                        ["Royalty Rate", "70%", "Optimal pricing threshold INR 99-299"],
                        ["Market Growth", "32% YoY", "Vernacular titles outperforming"],
                    ]
                },
                "callout": "Regional fiction exhibits highest retention."
            }
        ]
    }
    success = _build_fallback_document("docx", target_docx, content_plan)
    assert success is True
    assert target_docx.exists()
    assert target_docx.stat().st_size > 0

    doc = Document(str(target_docx))
    all_text = " ".join([p.text for p in doc.paragraphs])
    assert "Amazon KDP Business Opportunities in India" in all_text
    assert "Market Sizing and Vernacular Demand" in all_text
    assert len(doc.tables) == 1
    table_text = " ".join([cell.text for row in doc.tables[0].rows for cell in row.cells])
    assert "Royalty Rate" in table_text
    assert "70%" in table_text


def test_build_fallback_document_xlsx(tmp_path: Path):
    target_xlsx = tmp_path / "marketing_expenses.xlsx"
    content_plan = {
        "title": "Marketing Campaign Budget & Projections",
        "data_table": {
            "sheet_name": "Q3 Budget",
            "headers": ["Expense Item", "Allocated Budget (USD)", "Actual Spend", "Variance"],
            "rows": [
                ["Search Engine Ads", 5000, 4800, -200],
                ["Content Syndication", 3000, 3100, 100],
                ["Influencer Marketing", 4500, 4200, -300],
            ]
        }
    }
    success = _build_fallback_document("xlsx", target_xlsx, content_plan)
    assert success is True
    assert target_xlsx.exists()
    assert target_xlsx.stat().st_size > 0

    wb = openpyxl.load_workbook(str(target_xlsx))
    sheet = wb.active
    assert sheet.title == "Q3 Budget"
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["Expense Item", "Allocated Budget (USD)", "Actual Spend", "Variance"]
    row2 = [cell.value for cell in sheet[2]]
    assert row2[0] == "Search Engine Ads"
    assert row2[1] == 5000


def test_build_fallback_document_pptx(tmp_path: Path):
    target_pptx = tmp_path / "presentation.pptx"
    content_plan = {
        "title": "Shield AI Sovereign Enterprise",
        "subtitle": "Security and Agentic Architecture",
        "sections": [
            {
                "heading": "1. Sovereign Security Architecture",
                "paragraphs": ["Complete on-premise AST sandbox execution boundary."],
                "bullet_points": ["Zero external egress without verification.", "Automated policy compliance."],
            }
        ]
    }
    success = _build_fallback_document("pptx", target_pptx, content_plan)
    assert success is True
    assert target_pptx.exists()
    assert target_pptx.stat().st_size > 0

    prs = Presentation(str(target_pptx))
    assert len(prs.slides) == 2
    slide1_text = " ".join([shape.text for shape in prs.slides[0].shapes if shape.has_text_frame])
    assert "Shield AI Sovereign Enterprise" in slide1_text


def test_plan_and_generate_content_structure():
    import asyncio
    query = "Create a report on Amazon KDP business opportunities in India"
    plan = asyncio.run(plan_and_generate_content(query, doc_format="docx"))
    assert plan is not None
    assert "topic" in plan
    assert "title" in plan
    assert "scope_lock" in plan
    assert len(plan["sections"]) >= 1
    assert any("kdp" in s["heading"].lower() or "market" in s["heading"].lower() or "amazon" in s["heading"].lower() or "summary" in s["heading"].lower() for s in plan["sections"])
